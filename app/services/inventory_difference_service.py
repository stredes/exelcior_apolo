from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.db.database import (
    delete_inventory_difference,
    list_inventory_differences,
    upsert_inventory_difference,
)
from app.utils.utils import autoajustar_columnas


DIFFERENCE_COLUMNS = [
    "id",
    "Código",
    "Producto",
    "Bodega",
    "Ubicación",
    "N° Serie",
    "Lote",
    "Fecha Vencimiento",
    "Stock Sistema",
    "Dif. Stock",
    "Stock Contado",
    "Observación",
    "Archivo Origen",
    "Actualizado",
]


def build_inventory_item_key(record: Dict[str, object] | pd.Series) -> tuple[str, str, str, str, str]:
    if isinstance(record, pd.Series):
        data = record.to_dict()
    else:
        data = dict(record or {})
    return (
        str(data.get("Código", "") or "").strip(),
        str(data.get("Bodega", "") or "").strip(),
        str(data.get("Ubicación", "") or "").strip(),
        str(data.get("N° Serie", "") or "").strip(),
        str(data.get("Lote", "") or "").strip(),
    )


def save_inventory_difference(
    record: Dict[str, object] | pd.Series,
    *,
    difference_qty: int,
    note: str = "",
    source_file: str = "",
):
    if isinstance(record, pd.Series):
        data = record.to_dict()
    else:
        data = dict(record or {})

    stock_sistema = int(pd.to_numeric(pd.Series([data.get("Saldo Stock", 0)]), errors="coerce").fillna(0).iloc[0])
    return upsert_inventory_difference(
        codigo=str(data.get("Código", "") or "").strip(),
        producto=str(data.get("Producto", "") or "").strip(),
        bodega=str(data.get("Bodega", "") or "").strip(),
        ubicacion=str(data.get("Ubicación", "") or "").strip(),
        numero_serie=str(data.get("N° Serie", "") or "").strip(),
        lote=str(data.get("Lote", "") or "").strip(),
        fecha_vencimiento=str(data.get("Fecha Vencimiento", "") or "").strip(),
        stock_sistema=stock_sistema,
        diferencia_cantidad=int(difference_qty),
        observacion=str(note or "").strip(),
        archivo_origen=str(source_file or "").strip(),
    )


def get_inventory_differences_df() -> pd.DataFrame:
    registros = list_inventory_differences()
    rows: List[Dict[str, object]] = []
    for item in registros:
        actualizado = getattr(item, "actualizado_en", None) or getattr(item, "creado_en", None)
        rows.append(
            {
                "id": item.id,
                "Código": item.codigo or "",
                "Producto": item.producto or "",
                "Bodega": item.bodega or "",
                "Ubicación": item.ubicacion or "",
                "N° Serie": item.numero_serie or "",
                "Lote": item.lote or "",
                "Fecha Vencimiento": item.fecha_vencimiento or "",
                "Stock Sistema": int(item.stock_sistema or 0),
                "Dif. Stock": int(item.diferencia_cantidad or 0),
                "Stock Contado": int(item.stock_contado or 0),
                "Observación": item.observacion or "",
                "Archivo Origen": item.archivo_origen or "",
                "Actualizado": actualizado.strftime("%d/%m/%Y %H:%M") if actualizado else "",
            }
        )
    return pd.DataFrame(rows, columns=DIFFERENCE_COLUMNS)


def merge_inventory_differences(df: pd.DataFrame, diff_df: pd.DataFrame | None = None) -> pd.DataFrame:
    base = df.copy()
    if base.empty:
        if "Dif. Stock" not in base.columns:
            base["Dif. Stock"] = pd.Series(dtype="int64")
        if "Stock Contado" not in base.columns:
            base["Stock Contado"] = pd.Series(dtype="int64")
        if "Observación Dif." not in base.columns:
            base["Observación Dif."] = pd.Series(dtype="object")
        return base

    if diff_df is None:
        diff_df = get_inventory_differences_df()

    if diff_df is None or diff_df.empty:
        base["Dif. Stock"] = 0
        base["Stock Contado"] = pd.to_numeric(base["Saldo Stock"], errors="coerce").fillna(0).astype(int)
        base["Observación Dif."] = ""
        return base

    work = diff_df.copy()
    merge_keys = ["Código", "Bodega", "Ubicación", "N° Serie", "Lote"]
    work = work.loc[:, merge_keys + ["Dif. Stock", "Stock Contado", "Observación"]].rename(
        columns={"Observación": "Observación Dif."}
    )

    merged = base.merge(work, on=merge_keys, how="left")
    merged["Dif. Stock"] = pd.to_numeric(merged["Dif. Stock"], errors="coerce").fillna(0).astype(int)
    merged["Stock Contado"] = pd.to_numeric(merged["Stock Contado"], errors="coerce")
    merged["Stock Contado"] = merged["Stock Contado"].fillna(
        pd.to_numeric(merged["Saldo Stock"], errors="coerce").fillna(0)
    ).astype(int)
    merged["Observación Dif."] = merged["Observación Dif."].fillna("").astype(str)
    return merged


def remove_inventory_difference(record_id: int) -> bool:
    return delete_inventory_difference(record_id)


def build_inventory_difference_report_df(diff_df: pd.DataFrame | None = None) -> pd.DataFrame:
    if diff_df is None:
        diff_df = get_inventory_differences_df()

    if diff_df is None or diff_df.empty:
        return pd.DataFrame(columns=DIFFERENCE_COLUMNS[1:])

    report_columns = [
        "Código",
        "Producto",
        "Bodega",
        "Ubicación",
        "N° Serie",
        "Lote",
        "Fecha Vencimiento",
        "Stock Sistema",
        "Dif. Stock",
        "Stock Contado",
        "Observación",
        "Archivo Origen",
        "Actualizado",
    ]
    report = diff_df.loc[:, report_columns].copy()
    report["Dif. Stock"] = pd.to_numeric(report["Dif. Stock"], errors="coerce").fillna(0).astype(int)
    report["Stock Sistema"] = pd.to_numeric(report["Stock Sistema"], errors="coerce").fillna(0).astype(int)
    report["Stock Contado"] = pd.to_numeric(report["Stock Contado"], errors="coerce").fillna(0).astype(int)
    return report.sort_values(
        by=["Bodega", "Ubicación", "Producto", "Lote", "N° Serie"],
        kind="mergesort",
    ).reset_index(drop=True)


def export_inventory_difference_report(destination: str | Path, diff_df: pd.DataFrame | None = None) -> Path:
    report_df = build_inventory_difference_report_df(diff_df)
    if report_df.empty:
        raise ValueError("No hay diferencias guardadas para generar el informe.")

    output_path = Path(destination)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        report_df.to_excel(writer, sheet_name="Diferencias", startrow=2, index=False)
        workbook = writer.book
        sheet = writer.sheets["Diferencias"]

        title = f"INFORME DE DIFERENCIAS DE STOCK - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(report_df.columns))
        title_cell = sheet.cell(row=1, column=1, value=title)
        title_cell.font = Font(name="Segoe UI", bold=True, size=12)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx in range(1, len(report_df.columns) + 1):
            header_cell = sheet.cell(row=3, column=col_idx)
            header_cell.font = Font(name="Segoe UI", bold=True, size=10)
            header_cell.fill = header_fill
            header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            header_cell.border = border

        for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border

        sheet.freeze_panes = "A4"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
        autoajustar_columnas(workbook, max_width=40)

    return output_path
